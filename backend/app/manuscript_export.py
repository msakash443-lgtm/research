"""Manuscript export (X.31.13): the saved sections as Markdown, DOCX or JSON, pure functions with no database access.

A section is the ProseMirror JSON of its latest saved version (schema in `app/ledger.py`). Exports carry the researcher's
text and citations only; authorship marks (`prov`) are the provenance report's job and are left out. A citation (`cite`
node) is shown as "(Surname Year)" when its source is verified; an unverified or missing source shows a plain marker and
is never listed in the references (plan rule 24).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any, Iterator

from docx import Document as WordDocument
from docx.text.paragraph import Paragraph

UNVERIFIED_LABEL = "[unverified source]"
UNKNOWN_LABEL = "[unknown source]"
REFERENCES_TITLE = "References"


@dataclass(frozen=True)
class CitedSource:
    """A project source as export needs it. `verified` False means it is cited but must not be listed or exported."""

    id: str
    title: str
    authors: tuple[str, ...]
    year: int | None
    venue: str | None
    doi: str | None
    verified: bool


@dataclass(frozen=True)
class Section:
    title: str
    content: dict[str, Any]


def cited_source_ids(content: dict[str, Any]) -> list[str]:
    """Source ids cited in a document, in order of first use."""
    seen: dict[str, None] = {}
    for node in _walk(content):
        if node.get("type") == "cite":
            source_id = (node.get("attrs") or {}).get("source_id")
            if source_id:
                seen.setdefault(source_id, None)
    return list(seen)


def citation_label(source: CitedSource | None) -> str:
    if source is None:
        return UNKNOWN_LABEL
    if not source.verified:
        return UNVERIFIED_LABEL
    surname = _surname(source.authors[0]) if source.authors else source.title
    return f"({surname} {source.year})" if source.year else f"({surname})"


def reference_line(source: CitedSource) -> str:
    authors = "; ".join(source.authors) if source.authors else "Anonymous"
    year = f" ({source.year})" if source.year else ""
    parts = [f"{authors}{year}. {source.title.rstrip('.')}."]
    if source.venue:
        parts.append(f"{source.venue.rstrip('.')}.")
    if source.doi:
        parts.append(f"https://doi.org/{source.doi}")
    return " ".join(parts)


def references(sections: list[Section], sources: dict[str, CitedSource]) -> list[CitedSource]:
    """The verified sources cited anywhere in the manuscript, sorted by label then title."""
    ids: dict[str, None] = {}
    for section in sections:
        for source_id in cited_source_ids(section.content):
            ids.setdefault(source_id, None)
    cited = [sources[i] for i in ids if i in sources and sources[i].verified]
    return sorted(cited, key=lambda s: (citation_label(s).lower(), s.title.lower()))


def to_markdown(title: str, sections: list[Section], sources: dict[str, CitedSource]) -> str:
    out: list[str] = [f"# {title}", ""]
    for section in sections:
        out += [f"## {section.title}", ""]
        out += _markdown_blocks(section.content.get("content") or [], sources, indent="", heading_offset=1)
    refs = references(sections, sources)
    if refs:
        out += [f"## {REFERENCES_TITLE}", ""]
        out += [f"- {reference_line(s)}" for s in refs] + [""]
    return "\n".join(out).rstrip("\n") + "\n"


def to_docx(title: str, sections: list[Section], sources: dict[str, CitedSource]) -> bytes:
    word = WordDocument()
    word.core_properties.title = title
    word.add_heading(title, level=0)
    for section in sections:
        word.add_heading(section.title, level=1)
        _docx_blocks(word, section.content.get("content") or [], sources, list_style=None, depth=0, heading_offset=1)
    refs = references(sections, sources)
    if refs:
        word.add_heading(REFERENCES_TITLE, level=1)
        for source in refs:
            word.add_paragraph(reference_line(source), style="List Bullet")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def to_json(title: str, sections: list[dict[str, Any]], sources: list[CitedSource]) -> dict[str, Any]:
    """`sections` are already shaped by the caller (id, title, version, content_json, content_text)."""
    return {
        "title": title,
        "sections": sections,
        "references": [
            {"id": s.id, "title": s.title, "authors": list(s.authors), "year": s.year, "venue": s.venue, "doi": s.doi}
            for s in sources
        ],
    }


# ---------------------------------------------------------------- helpers


def _walk(node: dict[str, Any]) -> Iterator[dict[str, Any]]:
    yield node
    for child in node.get("content") or []:
        yield from _walk(child)


def _surname(author: str) -> str:
    """"Smith, Jane" -> Smith; "Jane Smith" -> Smith."""
    if "," in author:
        return author.split(",")[0].strip() or author.strip()
    words = author.split()
    return words[-1] if words else author.strip()


def _text_marks(node: dict[str, Any]) -> tuple[bool, bool]:
    names = {m.get("type") for m in node.get("marks") or []}
    return "strong" in names, "em" in names


def _inline_text(node: dict[str, Any], sources: dict[str, CitedSource]) -> str:
    kind = node.get("type")
    if kind == "text":
        return node.get("text", "")
    if kind == "cite":
        return citation_label(sources.get((node.get("attrs") or {}).get("source_id", "")))
    if kind == "hard_break":
        return "\n"
    return ""


def _markdown_inline(nodes: list[dict[str, Any]], sources: dict[str, CitedSource]) -> str:
    text = ""
    for node in nodes:
        piece = _inline_text(node, sources)
        if node.get("type") == "hard_break":
            text += "  \n"
            continue
        strong, em = _text_marks(node)
        core = piece.strip()
        if core and strong:
            core = f"**{core}**"
        if core and em:
            core = f"*{core}*"
        # Spaces stay outside the markers: "* word *" is not emphasis in Markdown.
        text += piece[: len(piece) - len(piece.lstrip())] + core + piece[len(piece.rstrip()):] if piece.strip() else piece
    return text


def _markdown_blocks(
    blocks: list[dict[str, Any]], sources: dict[str, CitedSource], *, indent: str, heading_offset: int
) -> list[str]:
    out: list[str] = []
    for block in blocks:
        kind = block.get("type")
        children = block.get("content") or []
        if kind == "paragraph":
            out += [indent + _markdown_inline(children, sources), ""]
        elif kind == "heading":
            level = min(int((block.get("attrs") or {}).get("level", 2)) + heading_offset, 6)
            out += [f"{'#' * level} {_markdown_inline(children, sources)}", ""]
        elif kind == "blockquote":
            inner = _markdown_blocks(children, sources, indent="", heading_offset=heading_offset)
            out += [f"> {line}" if line else ">" for line in inner]
        elif kind in ("bullet_list", "ordered_list"):
            for number, item in enumerate(children, start=1):
                marker = "-" if kind == "bullet_list" else f"{number}."
                inner = _markdown_blocks(item.get("content") or [], sources, indent=indent + " " * (len(marker) + 1), heading_offset=heading_offset)
                if inner:
                    inner[0] = indent + marker + " " + inner[0].lstrip()
                out += [line for line in inner if line != ""]
            out.append("")
    return out


def _docx_inline(paragraph: Paragraph, nodes: list[dict[str, Any]], sources: dict[str, CitedSource]) -> None:
    for node in nodes:
        if node.get("type") == "hard_break":
            paragraph.add_run().add_break()
            continue
        text = _inline_text(node, sources)
        if not text:
            continue
        run = paragraph.add_run(text)
        strong, em = _text_marks(node)
        run.bold = strong or None
        run.italic = em or None


def _docx_blocks(
    word: Any,
    blocks: list[dict[str, Any]],
    sources: dict[str, CitedSource],
    *,
    list_style: str | None,
    depth: int,
    heading_offset: int,
) -> None:
    for block in blocks:
        kind = block.get("type")
        children = block.get("content") or []
        if kind == "paragraph":
            style = list_style
            if style and depth > 1:
                style = f"{style} {min(depth, 3)}"
            paragraph = word.add_paragraph(style=style)
            _docx_inline(paragraph, children, sources)
        elif kind == "heading":
            level = min(int((block.get("attrs") or {}).get("level", 2)) + heading_offset - 1, 9)
            paragraph = word.add_heading(level=level)
            _docx_inline(paragraph, children, sources)
        elif kind == "blockquote":
            for child in children:
                if child.get("type") == "paragraph":
                    _docx_inline(word.add_paragraph(style="Quote"), child.get("content") or [], sources)
                else:
                    _docx_blocks(word, [child], sources, list_style=list_style, depth=depth, heading_offset=heading_offset)
        elif kind in ("bullet_list", "ordered_list"):
            style = "List Bullet" if kind == "bullet_list" else "List Number"
            for item in children:
                _docx_blocks(word, item.get("content") or [], sources, list_style=style, depth=depth + 1, heading_offset=heading_offset)

