"""Manuscript export (X.31.13): Markdown, DOCX, JSON, BibTeX, and Overleaf package."""

import io
import json
import uuid
import zipfile
from types import SimpleNamespace
import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import manuscript_export as me
from app.database import SessionLocal
from app.main import app
from app.models import AnalysisPlan, AnalysisRun, AuditEvent, DatasetVersion, Document, DocumentVersion, Source


def _login(email):
    client = TestClient(app)
    user_id = client.post("/api/auth/development/login", json={"email": email, "display_name": "D"}).json()["id"]
    return client, user_id


def _prov(sid, user_id):
    return {"type": "prov", "attrs": {"sid": sid, "origin": "typed", "by": user_id, "sug": None, "note": None, "paste": None, "src": None}}


def _text(text, sid, user_id, *extra):
    return {"type": "text", "text": text, "marks": [_prov(sid, user_id), *extra]}


def _cite(source_id):
    return {"type": "cite", "attrs": {"source_id": source_id}}


def _stat(run_id, path):
    return {"type": "stat", "attrs": {"analysis_run_id": run_id, "path": path}}


def _run(project_id, output, user_id="u1"):
    with SessionLocal() as db:
        dataset = DatasetVersion(
            project_id=uuid.UUID(project_id), version_no=1, kind="raw", name="d.csv", content="a\n1\n", sha256="0" * 64,
            row_count=1, column_names=["a"], created_by=user_id,
        )
        db.add(dataset)
        db.flush()
        plan = AnalysisPlan(
            project_id=uuid.UUID(project_id), dataset_id=dataset.id, spec_json={"tests": [], "variables": [], "seed": 0},
            plan_sha256="1" * 64, environment_sha256="2" * 64, created_by=user_id,
        )
        db.add(plan)
        db.flush()
        run = AnalysisRun(
            project_id=uuid.UUID(project_id), analysis_plan_id=plan.id, dataset_id=dataset.id, plan_sha256=plan.plan_sha256,
            dataset_sha256=dataset.sha256, script_sha256="3" * 64, environment_sha256=plan.environment_sha256,
            image="python@sha256:" + "4" * 64, output_json=output, created_by=user_id,
        )
        db.add(run)
        db.commit()
        return str(run.id)


def _new_document(project_id, user_id, title, position):
    with SessionLocal() as db:
        document = Document(project_id=uuid.UUID(project_id), title=title, position=position, created_by=f"user:{user_id}")
        db.add(document)
        db.commit()
        return str(document.id)


def _save(client, project_id, document_id, content):
    response = client.post(
        f"/api/projects/{project_id}/documents/{document_id}/versions",
        json={"base_number": 0, "client_save_id": uuid.uuid4().hex, "client_session_id": "s", "content_json": content},
    )
    assert response.status_code == 201, response.text


def _source(project_id, **fields):
    with SessionLocal() as db:
        row = Source(project_id=uuid.UUID(project_id), **fields)
        db.add(row)
        db.commit()
        return str(row.id)


@pytest.fixture
def world():
    client, user_id = _login("writer.export@example.org")
    project_id = client.post("/api/projects", json={"title": "Remote Work"}).json()["id"]
    verified = _source(
        project_id, title="Wages after remote work", authors=["Smith, Jane", "Lee, Ann"], year=2021,
        venue="Labour Journal", doi="10.1/abc", metadata_verified=True,
    )
    unverified = _source(project_id, title="Unchecked claim", authors=["Doe, Joe"], year=2020, metadata_verified=True)
    intro = _new_document(project_id, user_id, "Introduction", 1)
    methods = _new_document(project_id, user_id, "Methods", 2)
    _new_document(project_id, user_id, "Never saved", 3)
    _save(client, project_id, intro, {"type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 2}, "content": [_text("Background", "h", user_id)]},
        {"type": "paragraph", "content": [
            _text("Wages ", "a", user_id), _text("rose", "b", user_id, {"type": "strong"}), _text(" after ", "c", user_id),
            _cite(verified), _text(" and fell", "d", user_id, {"type": "em"}),
        ]},
        {"type": "bullet_list", "content": [
            {"type": "list_item", "content": [{"type": "paragraph", "content": [_text("first point", "e", user_id)]}]},
            {"type": "list_item", "content": [{"type": "paragraph", "content": [_text("second point", "f", user_id)]}]},
        ]},
    ]})
    _save(client, project_id, methods, {"type": "doc", "content": [
        {"type": "paragraph", "content": [_text("We used panel data.", "g", user_id)]},
    ]})
    # `unverified` stays verified and uncited here; dedicated tests below use it (or flip it) to prove the export guard.
    return SimpleNamespace(client=client, project_id=project_id, user_id=user_id, verified=verified, unverified=unverified)


def _export(world, fmt):
    return world.client.get(f"/api/projects/{world.project_id}/manuscript/export?format={fmt}")


def test_markdown_has_sections_in_order_formatting_and_references(world):
    response = _export(world, "md")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert 'filename="manuscript.md"' in response.headers["content-disposition"]
    text = response.text
    assert text.startswith("# Remote Work\n")
    assert text.index("## Introduction") < text.index("## Methods")
    assert "### Background" in text
    assert "Wages **rose** after (Smith 2021) *and fell*" in text
    assert "- first point" in text and "- second point" in text
    assert "## References" in text
    assert "Smith, Jane; Lee, Ann (2021). Wages after remote work. Labour Journal. https://doi.org/10.1/abc" in text
    assert "Unchecked claim" not in text.split("## References")[1]
    assert "Never saved" not in text
    assert response.headers["x-unsaved-sections"] == "1"


def test_authorship_marks_do_not_appear_in_the_export(world):
    text = _export(world, "md").text

    assert "prov" not in text and world.user_id not in text


def test_docx_opens_and_carries_headings_runs_and_references(world):
    response = _export(world, "docx")

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    word = WordDocument(io.BytesIO(response.content))
    texts = [p.text for p in word.paragraphs]
    assert texts[0] == "Remote Work"
    assert "Introduction" in texts and "Methods" in texts and "References" in texts
    body = next(p for p in word.paragraphs if p.text.startswith("Wages"))
    assert body.text == "Wages rose after (Smith 2021) and fell"
    assert [r.text for r in body.runs if r.bold] == ["rose"]
    assert [r.text for r in body.runs if r.italic] == [" and fell"]
    assert any(p.style.name == "List Bullet" and p.text == "first point" for p in word.paragraphs)


def test_json_lists_saved_sections_and_verified_references(world):
    body = json.loads(_export(world, "json").text)

    assert [s["title"] for s in body["sections"]] == ["Introduction", "Methods"]
    assert body["sections"][0]["version"] == 1
    assert body["sections"][0]["content_json"]["type"] == "doc"
    assert [r["id"] for r in body["references"]] == [world.verified]


def test_bibtex_has_only_cited_verified_sources(world):
    _source(world.project_id, title="Verified but never cited", authors=["Roe, Rae"], year=2019, metadata_verified=True)

    response = _export(world, "bibtex")

    assert response.status_code == 200
    assert 'filename="manuscript.bib"' in response.headers["content-disposition"]
    assert "Wages after remote work" in response.text
    assert "Unchecked claim" not in response.text
    assert "never cited" not in response.text


def test_nothing_saved_exports_only_the_title_and_unset_statement_placeholders(world):
    client, user_id = _login("empty.export@example.org")
    project_id = client.post("/api/projects", json={"title": "Empty"}).json()["id"]
    _new_document(project_id, user_id, "Draft", 1)

    response = client.get(f"/api/projects/{project_id}/manuscript/export?format=md")

    assert response.status_code == 200
    assert response.text.startswith("# Empty\n\n## Statements\n")
    assert "## Introduction" not in response.text and "## References" not in response.text
    assert "[Data Availability Statement: to complete]" in response.text
    assert response.headers["x-unsaved-sections"] == "1"


def test_unknown_format_is_rejected(world):
    assert _export(world, "pdf").status_code == 422


def test_export_is_audited_with_format_and_counts(world):
    _export(world, "docx")

    with SessionLocal() as db:
        events = db.scalars(select(AuditEvent).where(AuditEvent.action == "manuscript.exported", AuditEvent.project_id == uuid.UUID(world.project_id))).all()
    assert [e.payload_json for e in events] == [{"format": "docx", "sections": 2, "references": 1, "unsaved_sections": 1}]


def test_export_is_blocked_by_a_citation_to_an_unverified_source(world):
    doc = _new_document(world.project_id, world.user_id, "Discussion", 4)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [_text("See ", "x", world.user_id), _cite(world.unverified)]}]}
    with SessionLocal() as db:  # the save API refuses an unverified citation, so the version is written directly
        db.get(Source, uuid.UUID(world.unverified)).metadata_verified = False
        db.add(DocumentVersion(document_id=uuid.UUID(doc), project_id=uuid.UUID(world.project_id), number=1, content_json=content,
                               content_text="See", content_hash="0" * 64, saved_by=world.user_id, client_save_id="unver"))
        db.commit()

    response = _export(world, "md")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["citations"] == [{
        "document_id": doc, "document_title": "Discussion", "kind": "unverified_citation",
        "subject": world.unverified, "message": "Cited source is not verified: Unchecked claim",
        "status": "unverifiable", "verification": "not_checked",
    }]
    with SessionLocal() as db:
        events = db.scalars(select(AuditEvent).where(AuditEvent.action == "manuscript.export_blocked",
                                                      AuditEvent.project_id == uuid.UUID(world.project_id))).all()
    assert [e.payload_json for e in events] == [{"format": "md", "citations": [
        {"document_id": doc, "kind": "unverified_citation", "subject": world.unverified}
    ]}]


def test_a_blocked_citation_reports_the_verifiers_verdict(world):
    doc = _new_document(world.project_id, world.user_id, "Discussion", 4)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [_text("See ", "x", world.user_id), _cite(world.unverified)]}]}
    with SessionLocal() as db:
        source = db.get(Source, uuid.UUID(world.unverified))
        source.metadata_verified = False
        source.verification = {"verdict": "mismatch", "reasons": ["The year differs"]}
        db.add(DocumentVersion(document_id=uuid.UUID(doc), project_id=uuid.UUID(world.project_id), number=1, content_json=content,
                               content_text="See", content_hash="0" * 64, saved_by=world.user_id, client_save_id="mism"))
        db.commit()

    entry = _export(world, "json").json()["detail"]["citations"][0]

    assert (entry["status"], entry["verification"]) == ("unverifiable", "mismatch")


def test_json_references_carry_a_verified_status(world):
    references = json.loads(_export(world, "json").text)["references"]

    assert [r["verification"] for r in references] == [{"status": "verified", "method": "human"}]


def test_export_is_blocked_by_a_citation_to_a_source_in_another_project(world):
    other_client, _ = _login("other.export@example.org")
    other_project = other_client.post("/api/projects", json={"title": "Other"}).json()["id"]
    foreign = _source(other_project, title="Foreign", authors=["Fox, Al"], year=2018, metadata_verified=True)
    doc = _new_document(world.project_id, world.user_id, "Discussion", 4)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [_text("See ", "x", world.user_id), _cite(foreign)]}]}
    with SessionLocal() as db:  # the save API refuses this citation, so the version is written directly
        db.add(DocumentVersion(document_id=uuid.UUID(doc), project_id=uuid.UUID(world.project_id), number=1, content_json=content,
                               content_text="See", content_hash="0" * 64, saved_by=world.user_id, client_save_id="foreign"))
        db.commit()

    response = _export(world, "md")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["citations"] == [{
        "document_id": doc, "document_title": "Discussion", "kind": "unknown_source",
        "subject": foreign, "message": "Cited source is not an active source of this project",
    }]


def test_a_clean_export_is_not_blocked(world):
    assert _export(world, "md").status_code == 200


def test_markdown_and_docx_render_a_resolved_inserted_number(world):
    run = _run(world.project_id, {"descriptives": {"age": {"mean": 35.5}}})
    doc = _new_document(world.project_id, world.user_id, "Results", 5)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [
        _text("Mean age was ", "x", world.user_id), _stat(run, "descriptives.age.mean"),
    ]}]}
    _save(world.client, world.project_id, doc, content)

    md = _export(world, "md").text
    assert "Mean age was 35.5" in md

    word = WordDocument(io.BytesIO(_export(world, "docx").content))
    body = next(p for p in word.paragraphs if p.text.startswith("Mean age"))
    assert body.text == "Mean age was 35.5"


def test_export_is_blocked_by_a_number_naming_an_unknown_analysis_run(world):
    doc = _new_document(world.project_id, world.user_id, "Results", 5)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [_stat("nope", "mean")]}]}
    with SessionLocal() as db:  # the save API refuses this, so the version is written directly
        db.add(DocumentVersion(document_id=uuid.UUID(doc), project_id=uuid.UUID(world.project_id), number=1, content_json=content,
                               content_text="", content_hash="0" * 64, saved_by=world.user_id, client_save_id="unknown-run"))
        db.commit()

    response = _export(world, "md")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["citations"] == [{
        "document_id": doc, "document_title": "Results", "kind": "unknown_run",
        "subject": "nope", "message": "Inserted number names an analysis run this project does not have",
    }]


def test_export_is_blocked_by_a_number_whose_path_does_not_resolve(world):
    run = _run(world.project_id, {"descriptives": {"age": {"mean": 35.5}}})
    doc = _new_document(world.project_id, world.user_id, "Results", 5)
    content = {"type": "doc", "content": [{"type": "paragraph", "content": [_stat(run, "descriptives.age.missing")]}]}
    with SessionLocal() as db:  # the save API refuses this, so the version is written directly
        db.add(DocumentVersion(document_id=uuid.UUID(doc), project_id=uuid.UUID(world.project_id), number=1, content_json=content,
                               content_text="", content_hash="0" * 64, saved_by=world.user_id, client_save_id="unresolved-path"))
        db.commit()

    response = _export(world, "md")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["citations"][0]["kind"] == "unresolved_number"
    assert detail["citations"][0]["subject"] == f"{run}:descriptives.age.missing"


def test_pure_labels():
    source = me.CitedSource("1", "T", ("Jane Smith",), None, None, None, True)
    unverified = me.CitedSource("2", "U", ("Doe, Joe",), None, None, None, False)
    assert me.citation_label(source) == "(Smith)"
    assert me.citation_label(unverified) == me.UNVERIFIED_LABEL
    assert me.citation_label(None) == me.UNKNOWN_LABEL


def test_latex_export_has_sections_citations_statements_and_references(world):
    response = _export(world, "latex")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/x-latex")
    text = response.text
    assert r"\section{Introduction}" in text and r"\section{Methods}" in text
    assert r"\textbf{rose}" in text
    assert "(Smith 2021)" in text  # no discipline set: default style, unchanged
    assert r"\section*{References}" in text
    assert "Smith, Jane; Lee, Ann (2021)" in text
    assert r"\section*{Statements}" in text
    assert "[Data Availability Statement: to complete]" in text


def test_apa_style_from_the_discipline_profile_changes_citation_and_reference_punctuation(world):
    profile = world.client.put(f"/api/projects/{world.project_id}/profile", json={"discipline": "economics"})
    assert profile.status_code == 200 and profile.json()["effective"]["citation_style"] == "apa"

    text = _export(world, "md").text

    assert "(Smith, 2021)" in text  # apa: comma before the year, unlike the default style
    assert "Smith, Jane, & Lee, Ann (2021)." in text.split("## References")[1]


def test_vancouver_override_numbers_citations_by_first_appearance(world):
    override = world.client.put(
        f"/api/projects/{world.project_id}/profile", json={"discipline": None, "overrides": {"citation_style": "vancouver"}}
    )
    assert override.status_code == 200

    text = _export(world, "md").text

    assert "[1]" in text
    refs = text.split("## References")[1]
    assert "- 1. Smith, Jane, Lee, Ann. Wages after remote work." in refs


def test_unknown_citation_style_falls_back_to_the_default_format(world):
    override = world.client.put(
        f"/api/projects/{world.project_id}/profile", json={"discipline": None, "overrides": {"citation_style": "harvard-nonexistent"}}
    )
    assert override.status_code == 200

    assert "(Smith 2021)" in _export(world, "md").text


def test_manuscript_statements_are_null_until_set_then_carried_verbatim_into_every_export(world):
    before = world.client.get(f"/api/projects/{world.project_id}/manuscript/statements")
    assert before.status_code == 200
    before_body = before.json()
    assert before_body["data_availability_statement"] is None
    assert before_body["conflict_of_interest_statement"] is None
    assert before_body["ethics_statement"] is None
    assert "No AI-assisted actions are recorded" in before_body["ai_use_disclosure"]

    put = world.client.put(
        f"/api/projects/{world.project_id}/manuscript/statements",
        json={"data_availability_statement": "Data available on request.", "ethics_statement": "Approved by the IRB."},
    )
    assert put.status_code == 200
    assert put.json()["data_availability_statement"] == "Data available on request."
    assert put.json()["conflict_of_interest_statement"] is None

    md = _export(world, "md").text
    assert "Data available on request." in md
    assert "Approved by the IRB." in md
    assert "[Conflict of Interest Statement: to complete]" in md
    assert "[Data Availability Statement: to complete]" not in md

    docx_bytes = _export(world, "docx").content
    word = WordDocument(io.BytesIO(docx_bytes))
    assert any(p.text == "Data available on request." for p in word.paragraphs)

    body = json.loads(_export(world, "json").text)
    assert body["statements"]["data_availability"] == "Data available on request."
    assert body["statements"]["conflict_of_interest"] is None
    assert body["statements"]["ethics"] == "Approved by the IRB."
    assert "No AI-assisted actions are recorded" in body["statements"]["ai_use_disclosure"]

    with SessionLocal() as db:
        events = db.scalars(select(AuditEvent).where(AuditEvent.action == "manuscript.statements_updated")).all()
    assert [e.payload_json["fields_set"] for e in events] == [["data_availability_statement", "ethics_statement"]]


def test_ai_use_disclosure_is_generated_from_audit_events_that_carry_a_model_id(world):
    from app import audit

    with SessionLocal() as db:
        audit.record(
            db, actor="agent:research_run", action="research_run.answered", project_id=uuid.UUID(world.project_id),
            model_id="gpt-test", prompt_version="evidence_synthesis@2",
        )
        db.commit()

    statements = world.client.get(f"/api/projects/{world.project_id}/manuscript/statements").json()
    disclosure = statements["ai_use_disclosure"]
    assert "research_run.answered" in disclosure
    assert "gpt-test" in disclosure
    assert "evidence_synthesis@2" in disclosure
    assert "AI is not listed as an author" in disclosure

    md = _export(world, "md").text
    assert "## AI-Use Disclosure" in md
    assert "gpt-test" in md


def test_credit_contributions_default_empty_then_appear_in_every_export(world):
    get_before = world.client.get(f"/api/projects/{world.project_id}/manuscript/credit")
    assert get_before.status_code == 200
    rows = get_before.json()
    assert len(rows) == 1
    member_row = rows[0]
    assert member_row["roles"] == []

    md_before = _export(world, "md").text
    assert me.NO_CONTRIBUTORS in md_before

    put = world.client.put(
        f"/api/projects/{world.project_id}/manuscript/credit/{member_row['member_id']}",
        json={"roles": ["Conceptualization", "Writing - original draft"]},
    )
    assert put.status_code == 200
    assert put.json()["roles"] == ["Conceptualization", "Writing - original draft"]

    after = world.client.get(f"/api/projects/{world.project_id}/manuscript/credit").json()
    assert after[0]["roles"] == ["Conceptualization", "Writing - original draft"]

    md = _export(world, "md").text
    assert "## Author Contributions (CRediT)" in md
    assert "Conceptualization, Writing - original draft" in md

    docx_bytes = _export(world, "docx").content
    word = WordDocument(io.BytesIO(docx_bytes))
    assert any("Conceptualization, Writing - original draft" in p.text for p in word.paragraphs)

    body = json.loads(_export(world, "json").text)
    assert body["credit"] == [{"name": member_row["name"], "roles": ["Conceptualization", "Writing - original draft"]}]


def test_an_unknown_credit_role_is_refused_not_silently_dropped(world):
    rows = world.client.get(f"/api/projects/{world.project_id}/manuscript/credit").json()
    member_id = rows[0]["member_id"]
    response = world.client.put(
        f"/api/projects/{world.project_id}/manuscript/credit/{member_id}", json={"roles": ["Peer review"]},
    )
    assert response.status_code == 422
def test_overleaf_export_is_a_deterministic_zip_with_latex_and_verified_bibtex(world):
    first = _export(world, "overleaf")
    second = _export(world, "overleaf")
    assert first.status_code == 200
    assert first.headers["content-type"] == "application/zip"
    assert 'filename="overleaf-package.zip"' in first.headers["content-disposition"]
    assert first.content == second.content
    with zipfile.ZipFile(io.BytesIO(first.content)) as archive:
        assert set(archive.namelist()) == {"README.md", "manuscript.tex", "references.bib"}
        latex = archive.read("manuscript.tex").decode()
        bib = archive.read("references.bib").decode()
        assert "\\documentclass" in latex and "\\section{Introduction}" in latex
        assert "Wages after remote work" in bib
        assert "Unchecked claim" not in bib
        assert "Upload this ZIP to Overleaf" in archive.read("README.md").decode()
